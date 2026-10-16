import runpy
import unittest
from pathlib import Path
import openpyxl

module = runpy.run_path(str(Path(__file__).with_name('convert-lainsaadanto.py')))
assign_ids, kind, convert = module['assign_ids'], module['kind'], module['convert']


class LegislationImportTests(unittest.TestCase):
    def setUp(self):
        self.existing = [dict(id='d001', paatos='Ensimmäinen päätös', aika='1900'),
                         dict(id='d002', paatos='Toinen päätös', aika='2000')]

    def test_reordering_legacy_rows_preserves_references(self):
        records = [{'Päätös tai ratkaisu': d['paatos'], 'Aika': d['aika']} for d in reversed(self.existing)]
        self.assertEqual([r['Päätöstunniste'] for r in assign_ids(records, self.existing)], ['d002', 'd001'])

    def test_inserting_a_row_allocates_after_existing_ids_not_row_position(self):
        records = [{'Päätös tai ratkaisu': 'Uusi', 'Aika': '1950'}] + [
            {'Päätös tai ratkaisu': d['paatos'], 'Aika': d['aika']} for d in self.existing]
        with self.assertRaises(ValueError):
            assign_ids([dict(records[0])], self.existing)
        self.assertEqual([r['Päätöstunniste'] for r in assign_ids(records, self.existing, True)], ['d003', 'd001', 'd002'])

    def test_explicit_ids_allow_title_edits_and_duplicates_fail(self):
        self.assertEqual(assign_ids([{'Päätöstunniste': 'd002', 'Päätös tai ratkaisu': 'Muokattu'}], self.existing)[0]['Päätöstunniste'], 'd002')
        with self.assertRaises(ValueError):
            assign_ids([{'Päätöstunniste': 'd001'}, {'Päätöstunniste': 'd001'}], self.existing)

    def test_unknown_and_neutral_are_never_positive(self):
        for value in [None, '', 'Ei arvioitu', 'Vaikutukset avoimia']:
            self.assertEqual(kind(value), 'arvioimaton')
        self.assertEqual(kind('Neutraali'), 'neutraali')
        self.assertEqual(kind('Ristiriitainen, turvaa heikentävä ehto'), 'ristiriitainen')
        with self.assertRaises(ValueError):
            kind('Tuntematon uusi luokitus')

    def workbook(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = 'Sheet1'
        sheet.append(['Päätöstunniste', 'Aika', 'Päätös tai ratkaisu', 'Arvio'])
        for d in self.existing:
            sheet.append([d['id'], d['aika'], d['paatos'], 'Neutraali'])
        prep = wb.create_sheet('Valmistelu ja näyttö')
        prep.append(['Päätöstunniste', 'Päätös'])
        prep.append(['d002', 'Otsikko saa muuttua, kun tunniste on mukana'])
        wb.create_sheet('Rahoituskytkennät').append(['Päätös'])
        wb.create_sheet('Lukuohje').append(['Otsikko', 'Teksti'])
        return wb

    def test_supplementary_references_use_permanent_ids(self):
        output, _ = convert(self.workbook(), self.existing)
        self.assertEqual(output['valmistelu'][0]['paatosId'], 'd002')
        self.assertTrue(all(d['arvioLuokka'] == 'neutraali' for d in output['paatokset']))

    def test_deleted_decisions_or_broken_references_fail_before_output(self):
        wb = self.workbook()
        wb['Sheet1'].delete_rows(3)
        with self.assertRaises(ValueError):
            convert(wb, self.existing)

        wb = self.workbook()
        wb['Valmistelu ja näyttö']['A2'] = 'd999'
        with self.assertRaises(ValueError):
            convert(wb, self.existing)

    def test_unlinked_raw_rows_are_retained_separately_not_invented_as_references(self):
        wb = self.workbook()
        wb['Valmistelu ja näyttö'].append(['', 'Kohdentamaton raakapoiminta'])
        output, _ = convert(wb, self.existing)
        self.assertEqual(len(output['valmistelu']), 1)
        self.assertEqual(len(output['kohdentamatonValmistelu']), 1)
        self.assertEqual(output['kohdentamatonValmistelu'][0]['alkuperainenRivi']['Päätös'], 'Kohdentamaton raakapoiminta')


if __name__ == '__main__':
    unittest.main()
