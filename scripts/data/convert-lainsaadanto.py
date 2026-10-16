"""Excel -> legislation JSON, preserving permanent decision IDs.

First migrate an old workbook:
  python scripts/data/convert-lainsaadanto.py input.xlsx --write-ids with-ids.xlsx
Then keep the Päätöstunniste column when editing/reordering rows:
  python scripts/data/convert-lainsaadanto.py with-ids.xlsx
New decisions need a unique explicit ID (e.g. d117), or a migration pass to allocate one.
"""
import argparse
import json
from pathlib import Path
import re
import unicodedata
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
TARGET = ROOT / 'src/data/lainsaadanto.json'
ID_COLUMN = 'Päätöstunniste'
CLASSES = {'vahvistava', 'heikentava', 'korjaava', 'ristiriitainen', 'neutraali', 'arvioimaton'}


def text(value):
    return '' if value is None else str(value).strip()


def kind(value):
    normalized = unicodedata.normalize('NFKD', text(value).lower())
    normalized = ''.join(c for c in normalized if not unicodedata.combining(c))
    if normalized in {'', 'ei arvioitu', 'arvioimaton', 'vaikutukset avoimia',
                      'heikosti toteutettu murros', 'valttamattomyytta ja jakokysymyksia sisaltava',
                      'kestavyytta tavoitteleva, eriytyva'}:
        return 'arvioimaton'
    if normalized in {'neutraali', 'neutraali vaikutus'}:
        return 'neutraali'
    for stem, result in [('ristiriit', 'ristiriitainen'), ('korjaa', 'korjaava'),
                         ('heikent', 'heikentava'), ('vahvist', 'vahvistava')]:
        if stem in normalized:
            return result
    raise ValueError(f'Tuntematon arvio {value!r}; täsmennä arvio tai Arvioluokka.')


def rows(workbook, name):
    sheet = workbook[name]
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)]


def assign_ids(records, existing, allocate=False):
    by_key = {}
    for decision in existing:
        key = (text(decision['paatos']), text(decision['aika']))
        by_key.setdefault(key, []).append(decision['id'])
    used = set()
    reserved = {text(r.get(ID_COLUMN)) for r in records if text(r.get(ID_COLUMN))}
    next_number = max([int(d['id'][1:]) for d in existing if re.fullmatch(r'd\d+', d['id'])] +
                      [int(i[1:]) for i in reserved if re.fullmatch(r'd\d+', i)] + [0]) + 1
    for row in records:
        identifier = text(row.get(ID_COLUMN))
        if not identifier:
            candidates = by_key.get((text(row.get('Päätös tai ratkaisu')), text(row.get('Aika'))), [])
            if len(candidates) == 1:
                identifier = candidates[0]
            elif allocate:
                identifier = f'd{next_number:03d}'
                next_number += 1
            else:
                raise ValueError(f'Pysyvä {ID_COLUMN} puuttuu: {row.get("Päätös tai ratkaisu")!r}. Käytä --write-ids tai täytä tunniste.')
        if not re.fullmatch(r'd\d{3,}', identifier) or identifier in used:
            raise ValueError(f'Virheellinen tai päällekkäinen päätöstunniste: {identifier}')
        used.add(identifier)
        row[ID_COLUMN] = identifier
    return records


def convert(workbook, existing, allocate=False, allow_removals=False):
    records = assign_ids(rows(workbook, 'Sheet1'), existing, allocate)
    decisions = []
    for row in records:
        year = row.get('Päätösvuosi / jakson alku')
        try:
            year = int(float(year))
        except (TypeError, ValueError):
            match = re.search(r'\d{4}', text(row.get('Aika')))
            if not match:
                raise ValueError(f'Päätöksen vuosi puuttuu: {row[ID_COLUMN]}')
            year = int(match.group())
        explicit_class = text(row.get('Arvioluokka'))
        if explicit_class and explicit_class not in CLASSES:
            raise ValueError(f'Tuntematon Arvioluokka: {explicit_class}')
        decision = {'id': row[ID_COLUMN], 'aika': text(row.get('Aika')), 'vuosi': year,
                    'arvioLuokka': explicit_class or kind(row.get('Arvio'))}
        fields = {'paatos': 'Päätös tai ratkaisu', 'arvio': 'Arvio', 'metatarina': 'Merkitys metatarinalle',
                  'toimeenpano': 'Toimeenpano', 'teema': 'Teema', 'tavoite': 'Tavoite',
                  'riskinjako': 'Riskin ja vastuun jako', 'rajat': 'Tulkinnan rajat', 'lahde': 'Lähde',
                  'lahdelinkki': 'Lähdelinkki', 'saadostunnus': 'Säädöstunnus tietokannassa',
                  'heViitteet': 'HE-viitteet', 'vaikutusnaytto': 'Toteutuneiden vaikutusten näyttö'}
        decision.update({key: text(row.get(column)) for key, column in fields.items()})
        if not decision['paatos']:
            raise ValueError(f'Päätöksen kuvaus puuttuu: {decision["id"]}')
        decisions.append(decision)
    missing = {d['id'] for d in existing} - {d['id'] for d in decisions}
    if missing and not allow_removals:
        raise ValueError(f'Aiemmat tunnisteet puuttuvat: {sorted(missing)}. Tarkista tunnisteet; tarkoituksellinen poisto vaatii --allow-removals.')
    by_text = {}
    for decision in decisions:
        by_text.setdefault(decision['paatos'], []).append(decision['id'])
    ids = {d['id'] for d in decisions}

    def reference(row):
        explicit = text(row.get(ID_COLUMN))
        if explicit:
            if explicit not in ids:
                raise ValueError(f'Tuntematon päätösviittaus: {explicit}')
            return explicit
        candidates = by_text.get(text(row.get('Päätös')), [])
        if len(candidates) != 1:
            raise ValueError(f'Päätösviittaus puuttuu tai on moniselitteinen: {row.get("Päätös")!r}')
        return candidates[0]

    prep, unlinked = [], []
    for row in rows(workbook, 'Valmistelu ja näyttö'):
        try:
            identifier = reference(row)
        except ValueError:
            if text(row.get(ID_COLUMN)):
                raise
            unlinked.append({'syy': 'Päätösviittaus puuttuu tai ei ole yksikäsitteinen. Tarkista ennen kohdentamista.',
                             'alkuperainenRivi': {text(k): v for k, v in row.items() if k is not None}})
            continue
        prep.append({'paatosId': identifier, 'he': text(row.get('HE tai tutkimus')),
                     'otsikko': text(row.get('Otsikko / aineisto')), 'yhteys': text(row.get('Yhteyden luonne')),
                     'linkki': text(row.get('Lähdelinkki')), 'tulkinta': text(row.get('Näytön tulkinta'))})
    funding = [{'paatosId': reference(r), 'momentti': text(r.get('Momentti')),
                'nimi': text(r.get('Momentin nimi tietokannassa')), 'vuosi': r.get('Aineiston vuosi'),
                'summa': r.get('Momentin summa tietokannassa (€)'), 'varmuus': text(r.get('Tietokannan varmuusluokka'))}
               for r in rows(workbook, 'Rahoituskytkennät')]
    guide = [{'otsikko': text(r[0]), 'teksti': text(r[1])}
             for r in workbook['Lukuohje'].iter_rows(min_row=2, values_only=True) if r[0]]
    return {'schema_version': 1, 'lahde': 'Lainsaadanto.xlsx (fi_state_causal_map.db)',
            'paatokset': decisions, 'valmistelu': prep, 'kohdentamatonValmistelu': unlinked, 'rahoitus': funding, 'lukuohje': guide}, records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--write-ids', type=Path, help='Write a migrated workbook copy with permanent IDs')
    parser.add_argument('--output', type=Path, default=TARGET)
    parser.add_argument('--allow-removals', action='store_true')
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.workbook, data_only=True)
    existing = json.loads(TARGET.read_text())['paatokset']
    output, records = convert(workbook, existing, bool(args.write_ids), args.allow_removals)
    if args.write_ids:
        if args.write_ids.resolve() == args.workbook.resolve():
            raise ValueError('--write-ids requires a separate output workbook path')
        workbook = openpyxl.load_workbook(args.workbook)
        sheet = workbook['Sheet1']
        headers = [c.value for c in sheet[1]]
        column = headers.index(ID_COLUMN) + 1 if ID_COLUMN in headers else sheet.max_column + 1
        sheet.cell(1, column, ID_COLUMN)
        index = 0
        for row_number in range(2, sheet.max_row + 1):
            if any(sheet.cell(row_number, c).value is not None for c in range(1, sheet.max_column + 1)):
                sheet.cell(row_number, column, records[index][ID_COLUMN])
                index += 1
        workbook.save(args.write_ids)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=1) + '\n')
    print(f'Validated {len(output["paatokset"])} decisions with permanent IDs.')
    if output['kohdentamatonValmistelu']:
        print(f'VAROITUS: {len(output["kohdentamatonValmistelu"])} valmisteluriviä säilytettiin kohdentamattomina. Tarkista tunnisteet.')


if __name__ == '__main__':
    main()
